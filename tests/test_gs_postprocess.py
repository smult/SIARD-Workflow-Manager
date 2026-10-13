"""
Ghostscript-etterbehandling av PDF/A-filer fra LibreOffice (BLOB-konvertering):

  * «spør»: test på et utvalg, dialog med målt og beregnet besparelse
  * nei i dialogen → ingen filer endres
  * ja → alle filene erstattes av kontrollerte, mindre filer (tekst/sidetall
    likt, fortsatt PDF/A), og tableX.xml får riktig length/digest
  * uten dialog (ingen GUI) → ingenting gjøres
  * signert PDF røres ikke

Krever LibreOffice og Ghostscript (hoppes over uten). Testfilene er syntetiske.

Kjør:  python -X utf8 tests/test_gs_postprocess.py
"""
from __future__ import annotations

import hashlib
import io
import re
import sys
import tempfile
import threading
from pathlib import Path

_ROOT = Path(__file__).resolve().parent.parent
for p in (_ROOT, _ROOT / "tests"):
    if str(p) not in sys.path:
        sys.path.insert(0, str(p))


def _ok(cond: bool, msg: str) -> None:
    if not cond:
        raise AssertionError(msg)
    print(f"  ✓ {msg}")


def _docx(i: int) -> bytes:
    import docx
    d = docx.Document()
    d.add_heading(f"Vedtak nr. {i}", 1)
    for k in range(6):
        d.add_paragraph(f"Avsnitt {k}: Lorem ipsum dolor sit amet, æøå ÆØÅ — dokument {i}. " * 4)
    p = d.add_paragraph()
    p.add_run("Fet tekst ").bold = True
    p.add_run("og kursiv tekst.").italic = True
    b = io.BytesIO()
    d.save(b)
    return b.getvalue()


def _run(tmp: Path, name: str, answer, mode: str = "spør", n_files: int = 6):
    import settings
    settings.set_config = lambda k, v: None
    from test_failure_report import _make_siard
    from siard_workflow.core.context import WorkflowContext
    from siard_workflow.operations.blob_convert_operation import BlobConvertOperation
    root = tmp / name
    files = {f"record{i}.bin": _docx(i) for i in range(1, n_files + 1)}
    lob = _make_siard(root, files, list(files))
    asked: list = []
    logs: list = []
    meta = {"file_logger": None,
            "progress_cb": lambda ev, **k: logs.append(k.get("msg", "")) if ev == "log" else None,
            "stop_event": threading.Event(), "pause_event": threading.Event()}
    if answer is not None:
        meta["ask_gs_postprocess_cb"] = lambda est: (asked.append(est), answer)[1]
    ctx = WorkflowContext(siard_path=tmp / f"{name}.siard", metadata=meta)
    ctx.extracted_path = root
    op = BlobConvertOperation()
    op.params.update(gs_postprocess=mode, gs_sample_size=3)
    r = op.run(ctx)
    return r, lob, asked, logs, root


def _plain_lo_pdf(tmp: Path, i: int) -> bytes:
    """Vanlig PDF (ikke PDF/A) fra LibreOffice — som en PDF fra fagsystemet."""
    from siard_workflow.core import lo_runner as R
    from siard_workflow.core.libreoffice import find_libreoffice
    src = tmp / f"sys{i}.docx"
    src.write_bytes(_docx(100 + i))
    ok, err = R.convert(find_libreoffice("soffice"), src, tmp / f"sys{i}", "pdf",
                        tmp / "lo_prof", 120)
    assert ok, err
    return (tmp / f"sys{i}" / f"sys{i}.pdf").read_bytes()


def _image_pdf() -> bytes:
    """Ukomprimert PDF med tapsfritt (PNG/Flate) bilde og tekst."""
    import fitz
    from PIL import Image
    b = io.BytesIO()
    Image.new("RGB", (300, 200), (30, 120, 200)).save(b, "PNG")
    d = fitz.open()
    pg = d.new_page()
    pg.insert_text((72, 72), "Skannet vedlegg med bilde")
    for y in (100, 330, 560):                      # samme bilde tre ganger
        pg.insert_image(fitz.Rect(72, y, 372, y + 200), stream=b.getvalue())
    return d.tobytes(deflate=False, garbage=0)


def test_system_and_archive(tmp: Path) -> None:
    print("test_system_and_archive")
    import zipfile
    import settings
    settings.set_config = lambda k, v: None
    import fitz
    from test_failure_report import _make_siard
    from test_signed_pdf import _signed_pdf
    from siard_workflow.core.context import WorkflowContext
    from siard_workflow.core.msg_to_pdf import pdfa_identification
    from siard_workflow.operations.blob_convert_operation import BlobConvertOperation
    sys_pdf, img_pdf, signed = _plain_lo_pdf(tmp, 1), _image_pdf(), _signed_pdf()
    arc_pdf = _plain_lo_pdf(tmp, 2)
    zb = io.BytesIO()
    with zipfile.ZipFile(zb, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("brev.pdf", arc_pdf)
        z.writestr("signert.pdf", signed)
        z.writestr("notat.txt", "Tekst i arkivet")
    files = {"record1.bin": sys_pdf, "record2.bin": img_pdf, "record3.bin": signed,
             "record4.bin": zb.getvalue()}
    root = tmp / "sysarc"
    lob = _make_siard(root, files, [])
    rows = "".join(f'  <row>\n    <c1>{i}</c1>\n    <c2 file="{n}" length="1" digestType="MD5" '
                   f'digest="00"/>\n  </row>\n' for i, n in enumerate(files, 1))
    (root / "content/schema0/table0/table0.xml").write_text("<table>\n" + rows + "</table>\n", "utf-8")
    logs: list = []
    ctx = WorkflowContext(siard_path=tmp / "SysArc.siard", metadata={
        "file_logger": None,
        "progress_cb": lambda ev, **k: logs.append(k.get("msg", "")) if ev == "log" else None,
        "stop_event": threading.Event(), "pause_event": threading.Event()})
    ctx.extracted_path = root
    op = BlobConvertOperation()
    op.params.update(gs_postprocess="ja")
    r = op.run(ctx)
    gp = r.data.get("gs_postprocess") or {}
    print("    ", {k: gp.get(k) for k in ("files", "before", "after", "by_origin", "became_pdfa", "skipped")})
    got = {p.stem: p for p in lob.iterdir()}
    f1 = got["record1"]
    _ok(pdfa_identification(f1) and f1.stat().st_size < len(sys_pdf),
        f"systemets PDF (ikke PDF/A): mindre og nå PDF/A ({len(sys_pdf):,} → {f1.stat().st_size:,})")
    f2 = got["record2"]
    with fitz.open(str(f2)) as d:
        filters = {d.xref_get_key(x, "Filter")[1] for x in range(1, d.xref_length())
                   if d.xref_get_key(x, "Subtype")[1] == "/Image"}
    _ok("/DCTDecode" not in filters, f"tapsfritt bilde er fortsatt tapsfritt ({filters})")
    _ok(got["record3"].read_bytes() == signed, "signert PDF fra systemet er urørt")
    with zipfile.ZipFile(got["record4"]) as z:
        names = z.namelist()
        brev = z.read(next(n for n in names if n.endswith("brev.pdf")))
        sig = z.read(next(n for n in names if n.endswith("signert.pdf")))
        txt = z.read(next(n for n in names if n.endswith("notat.txt")))
    (tmp / "brev_ut.pdf").write_bytes(brev)
    _ok(len(brev) < len(arc_pdf) and pdfa_identification(tmp / "brev_ut.pdf"),
        f"PDF i arkiv: mindre og PDF/A ({len(arc_pdf):,} → {len(brev):,})")
    _ok(sig == signed and txt == b"Tekst i arkivet", "signert PDF og tekst i arkivet er urørt")
    _ok(gp.get("by_origin", {}).get("system", 0) >= 1 and gp.get("by_origin", {}).get("arkiv") == 1,
        f"telles per opprinnelse ({gp.get('by_origin')})")
    xml = (root / "content/schema0/table0/table0.xml").read_text("utf-8")
    bad = [p.name for p in lob.iterdir()
           if f'length="{p.stat().st_size}"' not in xml
           or hashlib.md5(p.read_bytes()).hexdigest().upper() not in xml.upper()]
    _ok(not bad, f"tableX.xml har riktig length/digest for alle filene ({bad or 'ok'})")


def test_allow_jpeg(tmp: Path) -> None:
    """Tilvalget «Tillat JPEG-omkoding av bilder»: standard tapsfritt; med
    tilvalget kan bilder bli JPEG (mindre fil), og visuell kontroll gjelder."""
    print("test_allow_jpeg")
    import fitz
    from PIL import Image
    from siard_workflow.core import ghostscript as G, pdf_postprocess as PP
    import math
    img = Image.new("RGB", (400, 300))
    # Fotolignende (jevne overganger) — Ghostscript velger JPEG bare for slike
    # bilder; ren støy beholdes tapsfritt også med tilvalget
    img.putdata([(int(120 + 60 * math.sin(x / 23) * math.cos(y / 17)), (x * y) % 200,
                  (x + 2 * y) % 255) for y in range(300) for x in range(400)])
    b = io.BytesIO()
    img.save(b, "PNG")
    d = fitz.open()
    pg = d.new_page()
    pg.insert_text((72, 72), "Fotografi")
    pg.insert_image(fitz.Rect(72, 100, 472, 400), stream=b.getvalue())
    src = tmp / "foto.pdf"
    d.save(str(src))
    gs = G.active_ghostscript()

    def filters(p: Path) -> set:
        with fitz.open(str(p)) as doc:
            return {doc.xref_get_key(x, "Filter")[1] for x in range(1, doc.xref_length())
                    if doc.xref_get_key(x, "Subtype")[1] == "/Image"}
    r1 = PP.optimize(PP.PdfItem(src), tmp / "tapsfri.pdf", gs, lossless=True)
    r2 = PP.optimize(PP.PdfItem(src), tmp / "jpeg.pdf", gs, lossless=False)
    if r1.ok:
        _ok("/DCTDecode" not in filters(tmp / "tapsfri.pdf"), "standard: bildet er ikke JPEG")
    else:
        _ok(True, f"standard: tapsfri versjon ikke mindre — originalen beholdes ({r1.reason})")
    _ok(r2.ok and "/DCTDecode" in filters(tmp / "jpeg.pdf"),
        f"tilvalg på: bildet kodet om til JPEG ({src.stat().st_size:,} → {r2.after:,} byte)")
    _ok(r2.after < (r1.after if r1.ok else r1.before), "tilvalget gir mindre fil enn tapsfritt")
    _ok(src.exists() and PP._visually_equal(src, tmp / "jpeg.pdf"),
        "visuell kontroll godtar JPEG-versjonen (originalen urørt)")


def main() -> int:
    from siard_workflow.core import ghostscript as G
    from siard_workflow.core.libreoffice import find_libreoffice
    from siard_workflow.core.msg_to_pdf import pdfa_identification
    if not (G.active_ghostscript() and find_libreoffice("soffice")):
        print("LibreOffice/Ghostscript mangler — test hoppet over")
        return 0
    with tempfile.TemporaryDirectory(prefix="gspost_") as td:
        tmp = Path(td)

        print("test_declined")
        r, lob, asked, logs, _ = _run(tmp, "nei", False)
        sizes_lo = {p.name: p.stat().st_size for p in lob.iterdir()}
        _ok(len(asked) == 1, "dialogen vises én gang etter testen")
        est = asked[0]
        _ok(est["n_sample"] == 3 and est["n_total"] == 6 and est["sample_before"] > 0
            and "est_saved" in est and est["est_seconds"] >= 0,
            f"utvalg 3 av 6, besparelse {est['sample_saving_pct']:.1f} %, "
            f"beregnet spart {est['est_saved']:,} byte")
        _ok(not r.data.get("gs_postprocess"), "nei → ingen etterbehandling registrert")
        _ok(any("valgt bort — ingen filer endret" in m for m in logs), "loggen sier at ingenting ble endret")

        print("test_accepted")
        r2, lob2, asked2, logs2, root2 = _run(tmp, "ja", True)
        gp = r2.data.get("gs_postprocess") or {}
        sizes_gs = {p.name: p.stat().st_size for p in lob2.iterdir()}
        _ok(gp.get("files", 0) >= 1 and gp["after"] < gp["before"],
            f"{gp.get('files')} filer etterbehandlet: {gp.get('before'):,} → {gp.get('after'):,} byte")
        _ok(sum(sizes_gs.values()) < sum(sizes_lo.values()),
            "filene i uttrekket er mindre enn uten etterbehandling")
        _ok(all(pdfa_identification(p) for p in lob2.iterdir()), "alle filer er fortsatt PDF/A")
        xml = (root2 / "content/schema0/table0/table0.xml").read_text("utf-8")
        okmeta = True
        for p in lob2.iterdir():
            m = re.search(rf'file="[^"]*{p.name}"[^>]*', xml)
            if m and 'length="' in m.group(0):
                okmeta &= f'length="{p.stat().st_size}"' in m.group(0)
            if m and 'digest="' in m.group(0):
                okmeta &= hashlib.md5(p.read_bytes()).hexdigest().upper() in m.group(0).upper() \
                    or hashlib.sha256(p.read_bytes()).hexdigest().upper() in m.group(0).upper()
        _ok(okmeta, "tableX.xml har length/digest for de etterbehandlede filene")
        _ok(any("Ghostscript-etterbehandling ferdig" in m for m in logs2), "loggen viser resultatet")

        print("test_no_gui")
        r3, _l, asked3, logs3, _ = _run(tmp, "uten_gui", None)
        _ok(not asked3 and not r3.data.get("gs_postprocess")
            and any("ingen dialog tilgjengelig" in m for m in logs3),
            "uten dialog: ingenting gjøres, og loggen sier hvorfor")

        test_system_and_archive(tmp)
        test_allow_jpeg(tmp)

        print("test_signed_untouched")
        from siard_workflow.core import pdf_postprocess as PP
        from test_signed_pdf import _signed_pdf
        f = tmp / "s.pdf"
        signed = _signed_pdf()
        f.write_bytes(signed)
        res = PP.optimize(PP.PdfItem(f), tmp / "s_out.pdf", G.active_ghostscript())
        _ok(not res.ok and "signert" in res.reason and f.read_bytes() == signed
            and not (tmp / "s_out.pdf").exists(),
            "signert PDF røres ikke av etterbehandlingen")
    print("\nAlle tester OK")
    import os
    os._exit(0)


if __name__ == "__main__":
    main()
